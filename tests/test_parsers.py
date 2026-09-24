import io
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from app.ingest.parsers import parse_upload
from tests.document_samples import sample_file


def test_txt():
    kind, text = parse_upload("a.md", "hello 世界".encode(), 10)
    assert kind == "text" and "hello" in text


def test_pdf():
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 700, "password=abc12345 project demo")
    c.save()
    kind, text = parse_upload("doc.pdf", buf.getvalue(), 10)
    assert kind == "pdf"
    assert "password=abc12345" in text


def test_reject_unknown():
    try:
        parse_upload("a.exe", b"xx", 10)
        assert False
    except ValueError:
        pass


@pytest.mark.parametrize('extension,kind', [('xlsx', 'excel'), ('xls', 'excel'), ('csv', 'csv'), ('docx', 'word'), ('pdf', 'pdf')])
def test_common_documents(extension, kind):
    parsed_kind, text = parse_upload(f'资料.{extension.upper()}', sample_file(extension), 10)
    assert parsed_kind == kind
    assert 'password' in text and 'DocImport9!' in text
    if extension in {'xlsx', 'xls'}:
        assert '工作表：账号清单' in text
        assert 'B2 password: DocImport9!' in text


def test_xlsx_hidden_sheets_formulas_and_blank_cells():
    book = Workbook()
    book.active.append(['日期', '金额', '说明'])
    book.active.append([None, '=1+2', '测试'])
    sheet = book.create_sheet('隐藏资料')
    sheet.sheet_state = 'hidden'
    sheet.append(['password', '备注'])
    sheet.append(['Hidden123!', '隐藏行也要扫描'])
    sheet.row_dimensions[2].hidden = True
    buf = io.BytesIO()
    book.save(buf)
    _, text = parse_upload('a.xlsx', buf.getvalue(), 10)
    assert 'B2 金额: =1+2' in text
    assert '隐藏资料' in text
    assert 'A2 password: Hidden123!' in text
    assert 'None' not in text


def test_word_order_headers_and_nested_table():
    doc = Document()
    doc.add_paragraph('开头')
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = '表格内容'
    cell.add_table(rows=1, cols=1).cell(0, 0).text = 'password=Nested123!'
    doc.add_paragraph('结尾')
    doc.sections[0].header.paragraphs[0].text = '页眉内容'
    doc.sections[0].footer.paragraphs[0].text = '页脚内容'
    buf = io.BytesIO()
    doc.save(buf)
    _, text = parse_upload('a.docx', buf.getvalue(), 10)
    assert text.index('开头') < text.index('表格内容') < text.index('结尾')
    assert 'password=Nested123!' in text
    assert '页眉内容' in text and '页脚内容' in text


@pytest.mark.parametrize('encoding', ['utf-8-sig', 'gb18030'])
@pytest.mark.parametrize('delimiter', [',', ';', '\t'])
def test_csv_encoding_and_delimiter(encoding, delimiter):
    data = delimiter.join(['服务', 'password']) + '\n' + delimiter.join(['服务甲', 'CsvSecret9!'])
    _, text = parse_upload('a.csv', data.encode(encoding), 10)
    assert '服务甲' in text
    assert 'B2 password: CsvSecret9!' in text


def test_csv_quoted_newlines_and_empty_columns():
    _, text = parse_upload('a.csv', b'name,password,note\na,Secret123,"line1\nline2"\nb,,ok', 10)
    assert 'B2 password: Secret123' in text
    assert 'line1\nline2' in text
    assert 'C3 note: ok' in text


@pytest.mark.parametrize('extension', ['xlsx', 'xls', 'docx', 'pdf'])
def test_corrupt_documents_return_safe_error(extension):
    with pytest.raises(ValueError, match='文件解析失败') as exc:
        parse_upload(f'a.{extension}', b'password=NeverEcho9!', 10)
    assert 'NeverEcho' not in str(exc.value)


def test_legacy_word_message():
    with pytest.raises(ValueError, match='另存为 .docx'):
        parse_upload('a.doc', b'legacy', 10)


def test_pdf_encrypted_and_image_only():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ValueError, match='扫描件暂不支持 OCR'):
        parse_upload('blank.pdf', buf.getvalue(), 10)
    writer.encrypt('password')
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ValueError, match='已加密'):
        parse_upload('locked.pdf', buf.getvalue(), 10)


@pytest.mark.parametrize('extension', ['xlsx', 'docx'])
def test_empty_office_document(extension):
    buf = io.BytesIO()
    (Workbook() if extension == 'xlsx' else Document()).save(buf)
    with pytest.raises(ValueError, match='无可提取文本'):
        parse_upload(f'empty.{extension}', buf.getvalue(), 10)


def test_upload_size_limit():
    with pytest.raises(ValueError, match='超过 0MB'):
        parse_upload('a.xlsx', sample_file('xlsx'), 0)


def test_zip_expansion_limit():
    buf = io.BytesIO()
    with ZipFile(buf, 'w', compression=ZIP_DEFLATED) as archive:
        archive.writestr('large.xml', b'a' * (11 * 1024 * 1024))
    with pytest.raises(ValueError, match='解压后过大'):
        parse_upload('a.xlsx', buf.getvalue(), 1)


def test_sparse_oversized_sheet():
    book = Workbook()
    book.active['XFD1048576'] = 'oversize'
    buf = io.BytesIO()
    book.save(buf)
    with pytest.raises(ValueError, match='表格过大'):
        parse_upload('a.xlsx', buf.getvalue(), 10)
