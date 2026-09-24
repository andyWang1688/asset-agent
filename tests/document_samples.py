"""纯内存合成文件；不使用用户资料，也不访问运行中的知识库。"""
import io

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
import xlwt


def sample_file(extension: str, secret: str = 'DocImport9!') -> bytes:
    buf = io.BytesIO()
    if extension == 'xlsx':
        book = Workbook()
        sheet = book.active
        sheet.title = '账号清单'
        sheet.append(['服务', 'password', '备注'])
        sheet.append(['示例服务', secret, '维护资料'])
        book.save(buf)
        book.close()
    elif extension == 'xls':
        book = xlwt.Workbook()
        sheet = book.add_sheet('账号清单')
        for row, values in enumerate([['服务', 'password', '备注'], ['示例服务', secret, '维护资料']]):
            for col, value in enumerate(values):
                sheet.write(row, col, value)
        book.save(buf)
    elif extension == 'docx':
        doc = Document()
        doc.add_paragraph('维护资料')
        table = doc.add_table(rows=2, cols=3)
        for row, values in zip(table.rows, [['服务', 'password', '备注'], ['示例服务', secret, '维护资料']]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.save(buf)
    elif extension == 'csv':
        return f'服务,password,备注\n示例服务,{secret},维护资料'.encode('utf-8-sig')
    elif extension == 'pdf':
        pdf = canvas.Canvas(buf)
        pdf.drawString(100, 700, f'password={secret} import test')
        pdf.save()
    else:
        raise AssertionError(extension)
    return buf.getvalue()
