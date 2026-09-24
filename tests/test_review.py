"""真实审查 API，临时目录 + 模型/保险柜替身；禁止读取用户资料。"""
import json
import pytest
from app import db
import app.main as main
from tests.test_maintenance_plan_flow import client, submit, SECRET, finish  # noqa: F401
from tests.document_samples import sample_file


def inspect(client, sid, body):
    return client.post(f'/api/pending/submissions/{sid}/review', json=body)


def test_review_is_local_only_and_no_cache(client):
    sid, body = submit(client)
    response = inspect(client, sid, body)
    assert response.status_code == 200, response.text
    assert response.headers['cache-control'] == 'no-store'
    document = response.json()['documents'][0]
    assert SECRET in document['text'] and SECRET not in document['preview']
    assert all(SECRET not in unit['preview'] for unit in document['units'])
    assert not client.review_model.calls and not client.review_vault.created
    assert not db.list_sources() and not db.list_tasks()
    assert SECRET not in str(db.get_report(db.get_submission(sid)['report_id']))
    settings = main.app.state.ctx.settings
    # Review uses encrypted pending payload, not the original-files directory.
    assert not list(settings.inbox_dir.iterdir())


@pytest.mark.parametrize('extension', ['xlsx', 'xls', 'csv', 'docx', 'pdf'])
def test_review_coordinates_match_extracted_content(client, extension):
    response = client.post('/api/ingest', data={'session_id': client.review_session}, files={'file': ('sample.' + extension, sample_file(extension))})
    assert response.status_code == 200, response.text
    view = response.json()
    body = {'session_id': client.review_session, 'decisions': {f['id']: f['suggested_action'] for f in view['findings']}}
    response = inspect(client, view['submission_id'], body)
    assert response.status_code == 200, response.text
    doc = response.json()['documents'][0]
    assert 'DocImport9!' not in doc['preview']
    for u in doc['units']:
        assert doc['text'][u['start']:u['end']] == u['text']
        assert 'DocImport9!' not in u['preview']
    if extension in ('xlsx', 'xls', 'csv'):
        assert doc['sheets']
        assert next(u for u in doc['units'] if u['row'] == 2 and u['col'] == 2)['text'] == 'DocImport9!'
    assert not client.review_model.calls


def test_manual_mark_name_and_value_reach_vault_only_after_confirmation(client):
    text = '联络暗号：青竹小屋'
    v = client.post('/api/ingest', data={'session_id': client.review_session, 'text': text}).json()
    sid = v['submission_id']; start = text.index('青竹')
    fid = f'manual:document:{start}:{len(text)}'
    body = {'session_id': client.review_session, 'decisions': {fid: 'store'}, 'manual': [{'source': 'document', 'start': start, 'end': len(text)}], 'edits': {fid: {'name': '联络暗号', 'description': '日常联络'}}}
    r = inspect(client, sid, body)
    assert r.status_code == 200, r.text
    assert '青竹小屋' not in r.json()['documents'][0]['preview']
    assert '联络暗号' in r.json()['documents'][0]['preview']
    assert not client.review_model.calls and not client.review_vault.created
    result = finish(client, sid, body)
    assert all('青竹小屋' not in call for call in client.review_model.calls)
    assert client.review_vault.created[0].name == '联络暗号'
    assert '青竹小屋' not in db.get_report(v['report_id'])['preview']
    assert inspect(client, sid, body).status_code == 400


def test_allow_and_invalid_manual_ranges(client):
    sid, body = submit(client)
    fid = next(iter(body['decisions']))
    r = inspect(client, sid, {**body, 'decisions': {fid: 'allow'}})
    assert r.status_code == 200 and SECRET in r.json()['documents'][0]['preview']
    r = inspect(client, sid, {**body, 'manual': [{'source': 'document', 'start': -1, 'end': 999999}]})
    assert r.status_code == 400 and SECRET not in r.text
    other = client.post('/api/chat/sessions', json={'mode': 'maintain'}).json()['session_id']
    assert inspect(client, sid, {**body, 'session_id': other}).status_code == 400
    assert not client.review_model.calls and not client.review_vault.created


def test_invalid_metadata_never_reaches_model(client):
    sid, body = submit(client)
    fid = next(iter(body['decisions']))
    r = inspect(client, sid, {**body, 'edits': {fid: {'name': SECRET}}})
    assert r.status_code == 400 and SECRET not in r.text
    assert not client.review_model.calls


def test_manual_instruction_mark_is_separate_and_draft_is_not_persisted(client):
    instruction = '整理到青竹小屋'
    v = client.post('/api/ingest', data={'session_id': client.review_session, 'text': instruction}, files={'file': ('sample.csv', sample_file('csv'))}).json()
    sid = v['submission_id']; start = instruction.index('青竹'); fid = f'manual:instruction:{start}:{len(instruction)}'
    body = {'session_id': client.review_session, 'decisions': {**{f['id']: f['suggested_action'] for f in v['findings']}, fid: 'store'}, 'manual': [{'source': 'instruction', 'start': start, 'end': len(instruction)}], 'edits': {fid: {'name': '资料归属'}}}
    before = dict(db.get_submission(sid))
    r = inspect(client, sid, body)
    assert r.status_code == 200, r.text
    doc = next(d for d in r.json()['documents'] if d['id'] == 'instruction')
    assert doc['text'] == instruction and '青竹小屋' not in doc['preview']
    assert dict(db.get_submission(sid)) == before
    assert not client.review_model.calls


@pytest.mark.parametrize('invalid', ['merge', 'duplicate_header', 'empty_header', 'ragged_csv', 'word_image', 'scan_pdf'])
def test_format_gate_rejects_before_pending_or_model(client, invalid):
    from io import BytesIO
    from openpyxl import Workbook
    from zipfile import ZipFile
    from pypdf import PdfWriter
    buf = BytesIO(); ext = 'xlsx'
    if invalid in ('merge', 'duplicate_header', 'empty_header'):
        book = Workbook(); sheet = book.active
        sheet.append(['服务', '服务'] if invalid == 'duplicate_header' else ['服务', None] if invalid == 'empty_header' else ['服务', '密码'])
        sheet.append(['example', 'secret'])
        if invalid == 'merge': sheet.merge_cells('A1:B1')
        book.save(buf)
    elif invalid == 'ragged_csv':
        ext = 'csv'; buf.write(b'name,password\nonly-one-cell')
    elif invalid == 'word_image':
        ext = 'docx'
        with ZipFile(buf, 'w') as z:
            z.writestr('word/document.xml', '<doc><drawing/></doc>')
    else:
        ext = 'pdf'; writer = PdfWriter(); writer.add_blank_page(width=100, height=100); writer.write(buf)
    r = client.post('/api/ingest', data={'session_id': client.review_session}, files={'file': ('unsupported.'+ext, buf.getvalue())})
    assert r.status_code == 400, r.text
    assert not db.list_submissions() and not db.list_tasks() and not db.list_sources()
    assert not client.review_model.calls and not client.review_vault.created


def test_entire_line_manual_mark_replaces_partial_detection(client):
    sid, body = submit(client)
    inspected = inspect(client, sid, body).json()
    unit = inspected['documents'][0]['units'][0]
    fid = f"manual:document:{unit['start']}:{unit['end']}"
    body['decisions'] = {fid: 'store'}
    body['manual'] = [{'source': 'document', 'start': unit['start'], 'end': unit['end']}]
    r = inspect(client, sid, body)
    assert r.status_code == 200, r.text
    assert len(r.json()['findings']) == 1
    assert r.json()['findings'][0]['name'] == '手动保护内容'
    assert SECRET not in r.json()['documents'][0]['preview']


def test_allow_one_repeated_value_keeps_other_occurrence_protected(client):
    value = 'review-person@example.test'
    text = f'第一行 {value}\n第二行 {value}'
    v = client.post('/api/ingest', data={'session_id': client.review_session, 'text': text}).json()
    body = {'session_id': client.review_session, 'decisions': {f['id']: 'redact' for f in v['findings']}}
    before = inspect(client, v['submission_id'], body).json()
    first = before['documents'][0]['units'][0]
    for fid in first['finding_ids']: body['decisions'][fid] = 'allow'
    response = inspect(client, v['submission_id'], body).json()['documents'][0]
    assert response['units'][0]['preview'] == first['text']
    assert value not in response['units'][1]['preview']
    assert response['preview'].count(value) == 1


@pytest.mark.parametrize('allowed_row', [2, 3])
def test_repeated_excel_cell_allow_matches_confirmed_raw(client, allowed_row):
    from io import BytesIO
    from pathlib import Path
    from openpyxl import Workbook
    value = 'review-person@example.test'
    book = Workbook(); sheet = book.active; sheet.title = '服务账号'
    sheet.append(['服务', '说明', '区域', '邮箱'])
    sheet.append(['资料库', '合成测试', '本机', value])
    sheet.append(['备份', '合成测试', '本机', value])
    other = book.create_sheet('个人资料'); other.append(['姓名', '邮箱']); other.append(['示例联系人', value])
    buf = BytesIO(); book.save(buf)
    v = client.post('/api/ingest', data={'session_id': client.review_session}, files={'file': ('review.xlsx', buf.getvalue())}).json()
    body = {'session_id': client.review_session, 'decisions': {f['id']: 'redact' for f in v['findings']}}
    doc = inspect(client, v['submission_id'], body).json()['documents'][0]
    selected = next(u for u in doc['units'] if u['sheet'] == 0 and u['row'] == allowed_row and u['col'] == 4)
    assert selected['finding_ids']
    for fid in selected['finding_ids']: body['decisions'][fid] = 'allow'
    doc = inspect(client, v['submission_id'], body).json()['documents'][0]
    assert next(u for u in doc['units'] if u['id'] == selected['id'])['preview'] == value
    assert all(value not in u['preview'] for u in doc['units'] if u['id'] != selected['id'])
    result = finish(client, v['submission_id'], body)
    assert db.get_report(v['report_id'])['preview'] == doc['preview']
    source = db.get_source(result['source_id']); raw = Path(source['path']).read_text()
    assert raw.endswith(doc['preview']) and raw.count(value) == 1
    spans = json.loads(source['allowed_spans'])
    assert len(spans) == 1 and raw[spans[0][0]:spans[0][1]] == value


def test_allow_in_attached_text_survives_confirmation(client):
    value = 'review-person@example.test'
    v = client.post('/api/ingest', data={'session_id': client.review_session, 'text': f'联系 {value}'}, files={'file': ('notes.txt', b'local notes')}).json()
    body = {'session_id': client.review_session, 'decisions': {f['id']: 'allow' for f in v['findings']}}
    assert body['decisions']
    result = finish(client, v['submission_id'], body)
    assert value in db.get_source(result['source_id'])['instruction']


@pytest.mark.parametrize('value', ['123', 'short-value', 'review-person@example.test'])
@pytest.mark.parametrize('allowed_index', [0, 1, 2])
def test_explicit_allow_preserves_only_its_span_and_unmarked_repeats_stay_masked(value, allowed_index):
    from app.security.entries import apply_entries, build_entries
    from app.security.rules import Finding
    text = ' | '.join([value] * 4)
    findings = [Finding(id=str(i), kind='pii', rule='manual', span=(i*(len(value)+3), i*(len(value)+3)+len(value)), confidence=1,
                        evidence='', suggested_action='redact', detector='manual', value=value) for i in range(3)]
    entries = build_entries(findings, {f.id: 'allow' if i == allowed_index else 'redact' for i, f in enumerate(findings)})
    result, allowed = apply_entries(text, entries)
    pieces = result.split(' | ')
    assert pieces[allowed_index] == value
    assert all(p == '[REDACTED:manual]' for i, p in enumerate(pieces[:3]) if i != allowed_index)
    assert len(allowed) == 1
    a, b = allowed[0]
    assert result[a:b] == value and a == sum(len(p)+3 for p in pieces[:allowed_index])
    # Existing short-value policy is unchanged: global fallback only masks eligible values.
    from app.security.redactor import should_mask_value
    assert pieces[3] == ('[REDACTED:manual]' if should_mask_value(value) else value)


def test_review_returns_exact_fragment_offsets_for_multiline_unicode_cell(client):
    text = '😀连接信息\n邮箱 one@example.test\n备用 two@example.test'
    v = client.post('/api/ingest', data={'session_id': client.review_session, 'text': text}).json()
    body = {'session_id': client.review_session, 'decisions': {f['id']: 'store' for f in v['findings']}}
    r = inspect(client, v['submission_id'], body).json()
    findings = {f['id']: f for f in r['findings']}
    for unit in r['documents'][0]['units']:
        assert len(unit['preview_spans']) == len(unit['finding_ids'])
        for span in unit['preview_spans']:
            assert unit['preview'][span['start']:span['end']] == findings[span['finding_id']]['private_ref']


def test_entropy_defaults_are_readable_unique_and_do_not_include_value_hash():
    from app.security.entries import build_entries
    from app.security.rules import Finding
    findings = [Finding(id=str(i), kind='unknown_suspect', rule='entropy_token', span=(i*20, i*20+12), confidence=.5, evidence='', suggested_action='store', detector='entropy', value=f'Fixture{i}AbcZ') for i in range(4)]
    names = [e.name for e in build_entries(findings, {f.id: 'store' for f in findings})]
    assert names == ['待确认内容', '待确认内容-2', '待确认内容-3', '待确认内容-4']


def test_four_fragments_in_one_excel_cell_have_distinct_preview_ranges(client):
    from io import BytesIO
    from openpyxl import Workbook
    book = Workbook(); book.active.append(['说明', '内容'])
    values = ['one@example.test', 'two@example.test', 'three@example.test', 'four@example.test']
    book.active.append(['合成资料', '😀连接信息\n' + '\n'.join(values)])
    buf = BytesIO(); book.save(buf)
    v = client.post('/api/ingest', data={'session_id': client.review_session}, files={'file': ('fragments.xlsx', buf.getvalue())}).json()
    body = {'session_id': client.review_session, 'decisions': {f['id']: 'store' for f in v['findings']}}
    r = inspect(client, v['submission_id'], body).json()
    unit = next(u for u in r['documents'][0]['units'] if u['row'] == 2 and u['col'] == 2)
    assert len(unit['preview_spans']) == 4
    second = unit['preview_spans'][1]['finding_id']
    body['decisions'][second] = 'allow'
    body['edits'] = {unit['preview_spans'][2]['finding_id']: {'name': '备用邮箱'}}
    r = inspect(client, v['submission_id'], body).json()
    unit = next(u for u in r['documents'][0]['units'] if u['row'] == 2 and u['col'] == 2)
    refs = {f['id']: f for f in r['findings']}
    for n, span in enumerate(unit['preview_spans']):
        value = unit['preview'][span['start']:span['end']]
        assert value == (values[n] if span['finding_id'] == second else refs[span['finding_id']]['private_ref'])
    assert '备用邮箱' in unit['preview']
    assert all(value not in unit['preview'] for i, value in enumerate(values) if i != 1)
