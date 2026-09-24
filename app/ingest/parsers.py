"""本地文档转文本；结果统一交给接收层安全扫描，不调用模型或执行公式。"""
import csv
import io
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
import xlrd


MAX_TABLE_CELLS = 200_000
MAX_TEXT_CHARS = 2_000_000
SUPPORTED = "支持 Markdown / TXT / PDF / Excel（.xlsx、.xls）/ CSV / Word（.docx）"


class _ParseError(ValueError):
    """仅承载程序固定的用户提示，禁止透传解析库中的文件内容。"""


def _value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _table_text(rows, cells: list | None = None, *, strict: bool = False) -> str:
    # 首个非空行作为列名，同时保留它自身；每个值保留坐标与列名，
    # 避免 password 列与下一行的密码拆散后绕过键值式安全规则。
    headers = None
    parts = []
    count = 0
    size = 0
    for row_no, row in enumerate(rows, 1):
        count += len(row)
        if row_no > MAX_TABLE_CELLS or count > MAX_TABLE_CELLS:
            raise _ParseError("表格过大，请拆分后上传")
        values = [_value(v) for v in row]
        if not any(values):
            continue
        if strict:
            if headers is None and (any(not v for v in values) or len(set(values)) != len(values)):
                raise _ParseError("表头需为单行、非空且不重复，请整理后上传")
            if headers is not None and len(values) != len(headers):
                raise _ParseError("表格各行列数不一致，请整理后上传")
        for col, value in enumerate(values, 1):
            if not value:
                continue
            label = headers[col - 1] if headers and col <= len(headers) else ""
            line = f"{get_column_letter(col)}{row_no} {label}: {value}"
            if cells is not None:
                start = size + len(line) - len(value)
                cells.append({"row": row_no, "col": col, "start": start, "end": start + len(value)})
            parts.append(line)
            size += len(line) + 1
            if size > MAX_TEXT_CHARS:
                raise _ParseError("提取文本过长，请拆分后上传")
        if headers is None:
            headers = values
    return "\n".join(parts)


def _excel(data: bytes, suffix: str, layout: list | None = None) -> str:
    parts = []
    if suffix == ".xlsx":
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book.worksheets:  # 包括隐藏工作表，不因可见性跳过安全检查
                if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_TABLE_CELLS:
                    raise _ParseError("表格过大，请拆分后上传")
                cells = []
                text = _table_text(sheet.iter_rows(values_only=True), cells, strict=True)
                if text:
                    prefix = f"## 工作表：{sheet.title}\n"
                    offset = sum(len(p) + 2 for p in parts)
                    if layout is not None:
                        layout.append({"name": sheet.title, "title_start": offset + 7,
                                       "title_end": offset + len(prefix) - 1,
                                       "cells": [{**c, "start": c["start"] + offset + len(prefix),
                                                  "end": c["end"] + offset + len(prefix)} for c in cells]})
                    parts.append(prefix + text)
        finally:
            book.close()
    else:
        book = xlrd.open_workbook(file_contents=data, on_demand=True, formatting_info=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows * sheet.ncols > MAX_TABLE_CELLS:
                    raise _ParseError("表格过大，请拆分后上传")

                def rows():
                    for row_no in range(sheet.nrows):
                        yield [
                            xlrd.xldate_as_datetime(cell.value, book.datemode)
                            if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                            for cell in sheet.row(row_no)
                        ]

                if sheet.merged_cells:
                    raise _ParseError("Excel 含合并单元格，请拆分为单行表头的规则表格")
                cells = []
                text = _table_text(rows(), cells, strict=True)
                if text:
                    prefix = f"## 工作表：{sheet.name}\n"
                    offset = sum(len(p) + 2 for p in parts)
                    if layout is not None:
                        layout.append({"name": sheet.name, "title_start": offset + 7,
                                       "title_end": offset + len(prefix) - 1,
                                       "cells": [{**c, "start": c["start"] + offset + len(prefix),
                                                  "end": c["end"] + offset + len(prefix)} for c in cells]})
                    parts.append(prefix + text)
        finally:
            book.release_resources()
    return "\n\n".join(parts)


def _word(data: bytes) -> str:
    doc = Document(io.BytesIO(data))

    def blocks(container):
        parts = []
        for block in container.iter_inner_content():
            if isinstance(block, Table):
                # 递归读取嵌套表格，不执行宏、外链或嵌入对象。
                rows = ([blocks(cell) for cell in row.cells] for row in block.rows)
                text = _table_text(rows)
                if text:
                    parts.append("[表格]\n" + text)
            elif block.text.strip():
                parts.append(block.text)
        return "\n".join(parts)

    parts = [blocks(doc)]
    for section in doc.sections:
        for name in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
            container = getattr(section, name)
            if not container.is_linked_to_previous:
                text = blocks(container)
                if text:
                    parts.append(f"[{name}]\n{text}")
    return "\n\n".join(parts)


def parse_upload(filename: str, data: bytes, max_mb: int, layout: list | None = None) -> tuple[str, str]:
    suffix = Path(filename or "").suffix.lower()
    if len(data) > max_mb * 1024 * 1024:
        raise ValueError(f"文件超过 {max_mb}MB 限制")
    if suffix == ".doc":
        raise ValueError("旧版 .doc 暂不支持，请另存为 .docx 后上传")
    if suffix not in {".md", ".txt", ".text", ".pdf", ".xlsx", ".xls", ".csv", ".docx"}:
        raise ValueError(SUPPORTED)
    try:
        if suffix in {".xlsx", ".docx"}:
            with ZipFile(io.BytesIO(data)) as archive:
                if sum(f.file_size for f in archive.infolist()) > max_mb * 10 * 1024 * 1024:
                    raise _ParseError("文件解压后过大，请拆分后上传")
                for name in archive.namelist():
                    if suffix == ".xlsx" and name.startswith("xl/worksheets/") and name.endswith(".xml"):
                        root = ElementTree.fromstring(archive.read(name))
                        if any(e.tag.rsplit("}", 1)[-1] in {"mergeCell", "drawing", "legacyDrawing"} for e in root.iter()):
                            raise _ParseError("Excel 含合并单元格、图片或图表，请整理为单行表头的纯数据表")
                    if suffix == ".docx" and name.startswith("word/") and name.endswith(".xml"):
                        root = ElementTree.fromstring(archive.read(name))
                        if any(e.tag.rsplit("}", 1)[-1] in {"drawing", "pict", "object", "txbxContent", "ins", "del"} for e in root.iter()):
                            raise _ParseError("Word 含图片、文本框、嵌入对象或未接受的修订，请整理为正文和表格后上传")
        if suffix in {".md", ".txt", ".text"}:
            kind, text = "text", data.decode("utf-8-sig", errors="replace")
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise _ParseError("PDF 已加密，请解密后上传")
            parts = []
            for number, page in enumerate(reader.pages, 1):
                value = page.extract_text() or ""
                if not value.strip() or list(page.images):
                    raise _ParseError("PDF 含扫描页、图片或无可提取文本的页面（扫描件暂不支持 OCR），请上传纯文本 PDF")
                if value.strip():
                    parts.append(f"## 第 {number} 页\n{value}")
            kind, text = "pdf", "\n\n".join(parts)
            if not text.strip():
                raise _ParseError("PDF 无可提取文本（扫描件暂不支持 OCR）")
        elif suffix in {".xlsx", ".xls"}:
            kind, text = "excel", _excel(data, suffix, layout)
        elif suffix == ".csv":
            try:
                decoded = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                decoded = data.decode("gb18030")
            try:
                dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            cells = []
            kind, text = "csv", _table_text(csv.reader(io.StringIO(decoded), dialect, strict=True), cells, strict=True)
            if layout is not None:
                layout.append({"name": "CSV", "cells": cells})
        else:
            kind, text = "word", _word(data)
        if not text.strip():
            raise _ParseError("文件无可提取文本，请检查内容（图片暂不支持 OCR）")
        if len(text) > MAX_TEXT_CHARS:
            raise _ParseError("提取文本过长，请拆分后上传")
        return kind, text.strip()
    except _ParseError as e:
        raise ValueError(str(e)) from None
    except Exception:
        # 解析库异常可能带原文、文件名或 XML，不能回传到浏览器或日志。
        raise ValueError("文件解析失败，请确认格式正确、文件未损坏且未加密") from None
